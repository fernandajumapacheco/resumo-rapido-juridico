"""Exportadores simples e legíveis para os relatórios gerados pela aplicação."""

from __future__ import annotations

import re
from pathlib import Path


def _blocos(markdown: str):
    for linha in markdown.splitlines():
        limpa = linha.strip()
        if not limpa:
            yield "espaco", ""
        elif limpa.startswith("# "):
            yield "titulo", limpa[2:]
        elif limpa.startswith("## "):
            yield "secao", limpa[3:]
        elif limpa.startswith("- "):
            yield "item", limpa[2:]
        elif limpa.startswith("> "):
            yield "aviso", limpa[2:]
        else:
            yield "texto", limpa


def _sem_markdown(texto: str) -> str:
    texto = re.sub(r"\*\*(.*?)\*\*", r"\1", texto)
    texto = re.sub(r"_(.*?)_", r"\1", texto)
    return texto


def criar_docx(markdown: str, destino: Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import RGBColor
    from docx.shared import Inches, Pt

    documento = Document()
    secao = documento.sections[0]
    secao.page_width = Inches(8.5)
    secao.page_height = Inches(11)
    secao.top_margin = secao.bottom_margin = Inches(0.8)
    secao.left_margin = secao.right_margin = Inches(0.9)
    estilos = documento.styles
    estilos["Normal"].font.name = "Arial"
    estilos["Normal"].font.size = Pt(11)
    estilos["Title"].font.name = "Arial"
    estilos["Title"].font.size = Pt(24)
    estilos["Title"].font.color.rgb = RGBColor(0, 0, 0)
    estilos["Title"].paragraph_format.space_after = Pt(18)
    titulo_ppr = estilos["Title"]._element.get_or_add_pPr()
    bordas = titulo_ppr.find(qn("w:pBdr"))
    if bordas is not None:
        titulo_ppr.remove(bordas)
    for nivel in ("Heading 1", "Heading 2"):
        estilos[nivel].font.name = "Arial"
        estilos[nivel].font.color.rgb = RGBColor(0, 0, 0)
        estilos[nivel].paragraph_format.space_before = Pt(12)
        estilos[nivel].paragraph_format.space_after = Pt(6)
        estilos[nivel].paragraph_format.keep_with_next = True
    for tipo, conteudo in _blocos(markdown):
        conteudo = _sem_markdown(conteudo)
        if tipo == "titulo":
            paragrafo = documento.add_paragraph(conteudo, style="Title")
            paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif tipo == "secao":
            documento.add_heading(conteudo, level=1)
        elif tipo == "item":
            documento.add_paragraph(conteudo, style="List Bullet")
        elif tipo == "aviso":
            paragrafo = documento.add_paragraph()
            trecho = paragrafo.add_run(conteudo)
            trecho.bold = True
        elif tipo == "texto":
            documento.add_paragraph(conteudo)
    destino.parent.mkdir(parents=True, exist_ok=True)
    documento.save(destino)
    return destino


def criar_pdf(markdown: str, destino: Path) -> Path:
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

    estilos = getSampleStyleSheet()
    estilos.add(ParagraphStyle(name="TituloRelatorio", parent=estilos["Title"], fontName="Helvetica-Bold", fontSize=20, leading=24, alignment=TA_CENTER, textColor="#000000", spaceAfter=18, keepWithNext=1))
    estilos.add(ParagraphStyle(name="SecaoRelatorio", parent=estilos["Heading1"], fontName="Helvetica-Bold", fontSize=14, leading=18, textColor="#000000", spaceBefore=13, spaceAfter=7, keepWithNext=1))
    estilos.add(ParagraphStyle(name="CorpoRelatorio", parent=estilos["BodyText"], fontName="Helvetica", fontSize=10.5, leading=15, textColor="#000000", spaceAfter=6))
    estilos.add(ParagraphStyle(name="AvisoRelatorio", parent=estilos["BodyText"], fontName="Helvetica-Bold", fontSize=10, leading=14, textColor="#000000", spaceBefore=4, spaceAfter=10))
    historia = []
    itens = []

    def descarregar_itens():
        nonlocal itens
        if itens:
            historia.append(ListFlowable([ListItem(Paragraph(item, estilos["CorpoRelatorio"])) for item in itens], bulletType="bullet", leftIndent=18, bulletFontName="Helvetica"))
            itens = []

    for tipo, conteudo in _blocos(markdown):
        conteudo = _sem_markdown(conteudo).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        if tipo == "item":
            itens.append(conteudo)
            continue
        descarregar_itens()
        if tipo == "titulo":
            historia.append(Paragraph(conteudo, estilos["TituloRelatorio"]))
        elif tipo == "secao":
            historia.append(Paragraph(conteudo, estilos["SecaoRelatorio"]))
        elif tipo == "aviso":
            historia.append(Paragraph(conteudo, estilos["AvisoRelatorio"]))
        elif tipo == "texto":
            historia.append(Paragraph(conteudo, estilos["CorpoRelatorio"]))
        elif tipo == "espaco":
            historia.append(Spacer(1, 3))
    descarregar_itens()

    destino.parent.mkdir(parents=True, exist_ok=True)
    documento = SimpleDocTemplate(str(destino), pagesize=letter, rightMargin=0.75 * inch, leftMargin=0.75 * inch, topMargin=0.7 * inch, bottomMargin=0.7 * inch, title="Relatório Jurídico")
    documento.build(historia)
    return destino


def criar_formatos(markdown: str, base: Path) -> dict[str, Path]:
    return {
        "docx": criar_docx(markdown, base.with_suffix(".docx")),
        "pdf": criar_pdf(markdown, base.with_suffix(".pdf")),
    }
